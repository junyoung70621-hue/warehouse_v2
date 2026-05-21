# utils/uploads.py — 엑셀 업로드/사용내역 공통 로직
import io
import re
import pandas as pd

EXCEL_COL_MAP = {
    "item_name":      "자재명",
    "quantity":       "수량",
    "rack_no":        "랙번호",
    "shelf":          "단",
    "box_no":         "박스번호",
    "category_large": "대분류",
    "category_mid":   "중분류",
    "category_small": "소분류",
    "item_location":  "지역",
    "location":       "자재위치",
    "erp_name":       "ERP품명",
    "erp_code":       "ERP코드",
    "repair_manager": "수리담당자명",
    "notes":          "비고",
}

USAGE_COL_MAP = {
    "item_name": "자재명",
    "erp_code":  "ERP코드",
    "quantity":  "사용수량",
    "reason":    "사용사유",
}


def make_excel_buffer(df_export, sheet_name="Sheet1") -> io.BytesIO:
    safe = re.sub(r'[\\/*?:\[\]]', '_', sheet_name)[:31] or "Sheet1"
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name=safe)
    buf.seek(0)
    return buf


def make_usage_template_buffer(df_usage, sheet_name="Sheet1") -> io.BytesIO:
    """사용내역 양식 엑셀: 표 테두리 + 사용수량 열 노란색 음영."""
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    safe = re.sub(r'[\\/*?:\[\]]', '_', sheet_name)[:31] or "Sheet1"
    buf  = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_usage.to_excel(writer, index=False, sheet_name=safe)
        ws = writer.sheets[safe]

        thin  = Side(style="thin",   color="AAAAAA")
        thick = Side(style="medium", color="888888")
        border_header = Border(top=thick, bottom=thick, left=thin, right=thin)
        border_cell   = Border(top=thin,  bottom=thin,  left=thin, right=thin)

        fill_header = PatternFill("solid", fgColor="2D2D2D")
        fill_qty    = PatternFill("solid", fgColor="FFF9C4")
        fill_qty_hd = PatternFill("solid", fgColor="F9A825")
        font_header = Font(bold=True, color="FFFFFF", size=10)
        font_qty_hd = Font(bold=True, color="FFFFFF", size=10)
        font_body   = Font(size=10)

        qty_col_idx = next(
            (i for i, c in enumerate(df_usage.columns, 1) if c == "사용수량"), None
        )

        n_cols = len(df_usage.columns)
        n_rows = len(df_usage) + 1
        for col_idx in range(1, n_cols + 1):
            is_qty = (col_idx == qty_col_idx)
            cell = ws.cell(row=1, column=col_idx)
            cell.fill      = fill_qty_hd if is_qty else fill_header
            cell.font      = font_qty_hd if is_qty else font_header
            cell.border    = border_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            for row_idx in range(2, n_rows + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.fill      = fill_qty if is_qty else PatternFill()
                c.font      = font_body
                c.border    = border_cell
                c.alignment = Alignment(
                    horizontal="center" if is_qty else "left", vertical="center"
                )
        col_widths = {"자재명": 35, "ERP코드": 16, "사용수량": 12, "사용사유": 22}
        for idx, col in enumerate(df_usage.columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = col_widths.get(col, 14)
        ws.row_dimensions[1].height = 18
    buf.seek(0)
    return buf


def clean_records(records: list) -> list:
    return [
        {k: (None if (v != v or str(v) in ("nan", "NaN", "None")) else v)
         for k, v in r.items()}
        for r in records
    ]


def validate_upload(up_df: pd.DataFrame, user_role: str, user: dict, centers: list) -> tuple:
    """엑셀 업로드 유효성 검사. (True, "") or (False, 오류메시지)"""
    if "location" not in up_df.columns:
        return True, ""
    invalid = set(up_df["location"].dropna().unique()) - set(centers)
    if invalid:
        return False, (
            f"❌ 등록되지 않은 센터명이 포함되어 있습니다: {', '.join(invalid)}\n\n"
            f"허용된 센터: {', '.join(centers)}"
        )
    if user_role != "admin":
        assigned  = user.get("assigned_center") or user.get("center", "")
        different = set(up_df["location"].dropna().unique()) - {assigned}
        if different:
            return False, (
                f"❌ 소속 센터({assigned}) 데이터만 업로드 가능합니다.\n\n"
                f"업로드 파일에 포함된 다른 센터: {', '.join(different)}"
            )
    return True, ""


def process_usage_upload(records: list, center: str, actor: dict):
    """사용내역 처리: 자재명/ERP코드로 매칭 후 수량 차감 + 이력 기록."""
    from utils.db import (
        get_supabase, clear_warehouse_cache,
        clear_history_cache, clear_usage_history_cache,
    )
    sb = get_supabase()
    ok, not_found, insufficient = 0, [], []

    for r in records:
        item_name = str(r.get("item_name", "") or "").strip()
        erp_code  = str(r.get("erp_code",  "") or "").strip()
        qty       = int(r.get("quantity", 0) or 0)
        reason    = str(r.get("reason", "사용내역 업로드") or "사용내역 업로드").strip()

        if (not item_name and not erp_code) or qty <= 0:
            continue

        item = None
        if erp_code:
            res = sb.table("warehouse").select("id,quantity") \
                    .eq("erp_code", erp_code).eq("location", center).execute()
            if res.data:
                item = res.data[0]
        if not item and item_name:
            res = sb.table("warehouse").select("id,quantity") \
                    .eq("item_name", item_name).eq("location", center).execute()
            if res.data:
                item = res.data[0]

        label = item_name or erp_code
        if not item:
            not_found.append(label)
            continue

        before = item["quantity"]
        if before < qty:
            insufficient.append(f"{label} (현재:{before} / 요청:{qty})")
            continue

        after = before - qty
        sb.table("warehouse").update({
            "quantity":         after,
            "last_modified_by": actor["id"],
            "last_modified_at": "now()",
        }).eq("id", item["id"]).execute()
        sb.table("history").insert({
            "actor_id":            actor["id"],
            "item_id":             item["id"],
            "action_type":         "out",
            "quantity":            qty,
            "reason":              reason,
            "from_center":         center,
            "snapshot_qty_before": before,
            "snapshot_qty_after":  after,
        }).execute()
        ok += 1

    clear_warehouse_cache()
    clear_history_cache()
    clear_usage_history_cache()
    return ok, not_found, insufficient
