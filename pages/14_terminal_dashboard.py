# pages/14_terminal_dashboard.py
import io
import uuid
from datetime import date, datetime, timedelta, timezone

_KST = timezone(timedelta(hours=9))

def _today_kst() -> date:
    return datetime.now(_KST).date()

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from utils.auth import is_role, require_login, logout
from utils.db import get_supabase
from utils.permissions import get_center as _get_center
from utils.routing import CENTERS
from utils.ui import (
    apply_global_css,
    render_sidebar_header,
    render_sidebar_section,
    render_sidebar_user,
    render_top_bar,
)

st.set_page_config(
    page_title="에이텍모빌리티 자재관리",
    page_icon="favicon_32.png",
    layout="wide",
    initial_sidebar_state="expanded",
)

apply_global_css()
require_login()

user        = st.session_state.user
user_role   = user["role"]
user_center = _get_center(user)

_is_admin   = user_role == "admin"
_is_jjae    = user_center == "자재센터"
# 출고 업로드: admin 또는 자재센터 소속 (비게스트)
_can_up_out = _is_admin or (_is_jjae and user_role != "guest")
# 입고 업로드: 게스트 제외 전원 (본인 센터 고정, admin/자재센터는 센터 선택 가능)
_can_up_in  = user_role != "guest"
# 입고 업로드 시 센터 선택 권한: admin 또는 자재센터
_can_sel_in = _is_admin or _is_jjae

NON_HUB_CENTERS = [c for c in CENTERS if c != "자재센터"]

TABLE = "terminal_movements"

# ── SQL 안내문 (테이블 미생성 시 표시) ────────────────────────────────────────
_SQL_SETUP = """\
-- Supabase SQL Editor에서 실행하세요
CREATE TABLE IF NOT EXISTS terminal_movements (
    id           uuid        DEFAULT gen_random_uuid() PRIMARY KEY,
    upload_id    uuid        NOT NULL,
    trcn_id      text        NOT NULL,
    device_type  text        NOT NULL,
    sub_type     text        NOT NULL,
    from_center  text        NOT NULL,
    to_center    text        NOT NULL,
    direction    text        NOT NULL CHECK (direction IN ('in','out')),
    uploaded_by  uuid        REFERENCES users(id) ON DELETE SET NULL,
    uploaded_at  timestamptz DEFAULT now(),
    upload_date  date        NOT NULL,
    file_name    text,
    notes        text
);
-- 기존 테이블에 컬럼 추가 (최초 1회만 실행)
ALTER TABLE terminal_movements ADD COLUMN IF NOT EXISTS notes text;
CREATE INDEX IF NOT EXISTS idx_tm_upload_date ON terminal_movements(upload_date);
CREATE INDEX IF NOT EXISTS idx_tm_direction   ON terminal_movements(direction);
CREATE INDEX IF NOT EXISTS idx_tm_trcn_id     ON terminal_movements(trcn_id);
"""


# ══════════════════════════════════════════════════════════════════════════════
# 단말기 분류 로직
# ══════════════════════════════════════════════════════════════════════════════

def classify_terminal(raw) -> tuple[str, str]:
    """TRCN_ID → (단말기종류, 유형).  숫자만 추출 후 길이·prefix 기준 분류."""
    # Excel이 숫자 셀을 float으로 읽으면 '560002215.0' 형태로 전달됨 → 정수 문자열로 정규화
    s = str(raw).strip()
    try:
        s = str(int(float(s)))
    except (ValueError, OverflowError):
        pass
    digits = "".join(c for c in s if c.isdigit())
    if not digits:
        return "미분류", "알 수 없음"

    n = len(digits)

    # ── 승하차·운전자 계열 (8~9자리) ──
    if n in (8, 9):
        if digits.startswith("157"):   return "한강버스", "승하차"
        if digits.startswith("1560"):  return "B800", "승하차"
        if digits.startswith("1553"):  return "B710", "승하차"
        if digits.startswith("1551"):  return "B620", "승하차"
        if digits.startswith("1451"):  return "B620", "운전자"

    # ── 표출기·통합단말기 계열 (9자리) ──
    if n == 9:
        if digits.startswith("5600"):  return "B800", "표출기"
        if digits.startswith("5500"):  return "B800", "통합단말기"
        if digits.startswith("457"):   return "한강버스", "표출기"
        if digits.startswith("447"):   return "한강버스", "통합단말기"
        if digits.startswith("4550"):  return "B710", "표출기"
        if digits.startswith("4450"):  return "B710", "통합단말기"
        if digits.startswith("4500"):  return "B700", "표출기"
        if digits.startswith("4400"):  return "B700", "통합단말기"

    # ── 6자리 모뎀 (더 구체적인 prefix 먼저) ──
    if n == 6:
        if digits.startswith("10"):
            if 100001 <= int(digits) <= 100500:               return "B620", "모뎀"
            else:                                             return "B800", "모뎀"
        if digits.startswith("6"):                             return "B710", "모뎀"
        if digits.startswith("4") or digits.startswith("5"):  return "B700", "모뎀"
        if digits.startswith("1"):                             return "B620", "모뎀"

    return "미분류", "알 수 없음"


DEVICE_ORDER = ["B800", "B700", "B710", "B620", "한강버스", "미분류"]
SUB_ORDER    = ["표출기", "통합단말기", "승하차", "운전자", "모뎀", "알 수 없음"]


def _apply_classifications(series: pd.Series) -> pd.DataFrame:
    """pd.Series → DataFrame[_dtype, _stype]"""
    results = [classify_terminal(v) for v in series]
    if results:
        return pd.DataFrame(results, columns=["_dtype", "_stype"], index=series.index)
    return pd.DataFrame({"_dtype": pd.Series(dtype=str), "_stype": pd.Series(dtype=str)})


def build_pivot(rows: list) -> pd.DataFrame:
    """행 목록 → 단말기종류 x 유형 pivot 테이블."""
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows)
    df["device_type"] = pd.Categorical(df["device_type"], categories=DEVICE_ORDER, ordered=True)
    df["sub_type"]    = pd.Categorical(df["sub_type"],    categories=SUB_ORDER,    ordered=True)
    pivot = df.pivot_table(
        index="device_type", columns="sub_type",
        values="trcn_id", aggfunc="count", fill_value=0, observed=False,
    )
    pivot.index.name   = "종류"
    pivot.columns.name = None
    present = [c for c in SUB_ORDER if c in pivot.columns]
    pivot   = pivot[present]
    pivot["합계"] = pivot.sum(axis=1)
    return pivot[pivot["합계"] > 0]


# 단말기 기종 표기 단축 매핑
_DEVICE_SHORT = {"B800": "B800", "B700": "B700", "B710": "B710", "B620": "B620", "한강버스": "한강버스", "미분류": "기타"}
_SUB_SHORT    = {"통합단말기": "통합", "표출기": "표출", "승하차": "승하차",
                 "운전자": "운전자", "모뎀": "모뎀", "알 수 없음": "기타"}
_ROW_ORDER = [
    "B800통합", "B800표출", "B800승하차", "B800모뎀",
    "B700통합", "B700표출", "B700승하차", "B700모뎀",
    "B710통합", "B710표출", "B710승하차", "B710모뎀",
    "B620승하차", "B620운전자", "B620모뎀",
    "한강버스통합", "한강버스표출", "한강버스승하차",
]
_CTR_SHORT = {
    "강남센터": "강남", "강동센터": "강동", "강북센터": "강북", "강서센터": "강서",
    "택시지원파트": "택시", "리페어팀": "리페어",
}
_COL_ORDER = ["강남", "강동", "강북", "강서", "택시", "리페어"]


def build_center_pivot(rows: list, direction: str = "out") -> pd.DataFrame | None:
    """단말기종류 × 센터 크로스표 (이미지 형식)."""
    if not rows:
        return None
    df = pd.DataFrame(rows)
    df["row_key"] = (
        df["device_type"].map(_DEVICE_SHORT).fillna(df["device_type"])
        + df["sub_type"].map(_SUB_SHORT).fillna(df["sub_type"])
    )
    ctr_col = "to_center" if direction == "out" else "from_center"
    df["ctr"] = df[ctr_col].map(_CTR_SHORT).fillna(df[ctr_col])

    pivot = df.pivot_table(
        index="row_key", columns="ctr",
        values="trcn_id", aggfunc="count", fill_value=0,
    )
    pivot.columns.name = None

    rows_ord  = [r for r in _ROW_ORDER if r in pivot.index]
    rows_ext  = [r for r in pivot.index if r not in _ROW_ORDER]
    cols_ord  = [c for c in _COL_ORDER if c in pivot.columns]
    cols_ext  = [c for c in pivot.columns if c not in _COL_ORDER]

    result = pivot.reindex(rows_ord + rows_ext)[cols_ord + cols_ext].fillna(0).astype(int)
    return result[result.sum(axis=1) > 0]


def render_center_table(pivot: pd.DataFrame, ref_date: date) -> None:
    """단말기종류×센터 크로스표를 Excel 이미지 스타일로 출력."""
    cols = list(pivot.columns)
    n    = len(cols)
    th_date = f"{ref_date.month}월 {ref_date.day}일"

    # ── 셀 스타일 상수 ──────────────────────────────────────────────
    S_TH_DATE  = "background:#FCE4ED;color:#D3004F;font-weight:700;text-align:center;padding:7px 10px;border:1px solid #f0c0d0;font-size:13px;"
    S_TH_CTR   = "background:#F8F9FA;color:#1E293B;font-weight:600;text-align:center;padding:5px 8px;border:1px solid #E2E8F0;font-size:12px;min-width:52px;"
    S_TH_LABEL = "background:#F8F9FA;color:#64748B;font-weight:600;text-align:center;padding:5px 8px;border:1px solid #E2E8F0;font-size:12px;min-width:80px;"
    S_ROW_HDR  = "background:#FEF3F6;color:#D3004F;font-weight:600;text-align:left;padding:5px 10px;border:1px solid #E2E8F0;font-size:12px;"
    S_CELL     = "text-align:center;padding:5px 8px;border:1px solid #E2E8F0;font-size:12px;color:#1E293B;"
    S_CELL_0   = "text-align:center;padding:5px 8px;border:1px solid #E2E8F0;font-size:12px;color:#CBD5E1;"

    col_heads = "".join(f"<th style='{S_TH_CTR}'>{c}</th>" for c in cols)
    rows_html = ""
    for row_key, row_data in pivot.iterrows():
        cells = ""
        for c in cols:
            v = row_data[c]
            cells += f"<td style='{S_CELL}'>{v}</td>" if v > 0 else f"<td style='{S_CELL_0}'></td>"
        rows_html += f"<tr><td style='{S_ROW_HDR}'>{row_key}</td>{cells}</tr>"

    html = f"""
    <div style="overflow-x:auto;">
    <table style="border-collapse:collapse;font-family:'Noto Sans KR',sans-serif;width:auto;">
      <thead>
        <tr><th colspan="{n+1}" style="{S_TH_DATE}">{th_date}</th></tr>
        <tr><th style="{S_TH_LABEL}">종류</th>{col_heads}</tr>
      </thead>
      <tbody>{rows_html}</tbody>
    </table>
    </div>
    """
    st.markdown(html, unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
# Excel 파싱
# ══════════════════════════════════════════════════════════════════════════════

def parse_terminal_excel(uploaded_file) -> pd.DataFrame | None:
    """헤더 행 자동 감지 후 파싱. xlsx→openpyxl, xls→xlrd."""
    fname = getattr(uploaded_file, "name", "")
    engine = "xlrd" if fname.lower().endswith(".xls") else "openpyxl"
    _kws = ["trcn_id", "trcnid", "단말기id", "단말기번호", "단말기 id",
            "단말기 번호", "ih번호", "단말기", "번호"]
    try:
        content = uploaded_file.read()
        raw = pd.read_excel(io.BytesIO(content), engine=engine, header=None, dtype=str).fillna("")
        header_row = 0
        for i, row in raw.iterrows():
            for val in row:
                v = str(val).strip().lower().replace(" ", "").replace("_", "")
                if any(kw in v for kw in _kws):
                    header_row = i
                    break
            else:
                continue
            break
        df = pd.read_excel(io.BytesIO(content), engine=engine, header=header_row, dtype=str)
        return df.dropna(how="all").reset_index(drop=True)
    except Exception as e:
        st.error(f"파일 읽기 실패: {e}")
        return None


def find_trcn_col(df: pd.DataFrame) -> str | None:
    """단말기ID 컬럼 자동 감지."""
    kws = ["trcn_id", "trcnid", "단말기id", "단말기번호", "단말기 id", "단말기 번호",
           "ih번호", "ih", "단말기", "번호"]
    for col in df.columns:
        nm = str(col).strip().lower().replace(" ", "").replace("_", "")
        for kw in kws:
            if kw in nm:
                return col
    return None


# ══════════════════════════════════════════════════════════════════════════════
# DB 함수
# ══════════════════════════════════════════════════════════════════════════════

def _table_exists() -> bool:
    try:
        get_supabase().table(TABLE).select("id").limit(1).execute()
        return True
    except Exception:
        return False


def fetch_terminal(
    direction: str = None,
    upload_date: date = None,
    date_from: date = None,
    date_to: date = None,
    limit: int = 3000,
) -> list:
    try:
        q = (
            get_supabase().table(TABLE)
            .select("id,upload_id,trcn_id,device_type,sub_type,"
                    "from_center,to_center,direction,"
                    "uploaded_at,upload_date,file_name,uploaded_by,notes")
            .order("uploaded_at", desc=True)
        )
        if direction:   q = q.eq("direction",   direction)
        if upload_date: q = q.eq("upload_date", upload_date.isoformat())
        if date_from:   q = q.gte("upload_date", date_from.isoformat())
        if date_to:     q = q.lte("upload_date", date_to.isoformat())
        return q.limit(limit).execute().data or []
    except Exception:
        return []


def check_dups(trcn_ids: list[str], upload_date: date, direction: str) -> set[str]:
    """이미 등록된 TRCN_ID 집합 반환 (같은 날짜·방향 기준)."""
    if not trcn_ids:
        return set()
    try:
        res = (
            get_supabase().table(TABLE)
            .select("trcn_id")
            .eq("upload_date", upload_date.isoformat())
            .eq("direction",   direction)
            .in_("trcn_id",    trcn_ids[:500])
            .execute()
        )
        return {r["trcn_id"] for r in (res.data or [])}
    except Exception:
        return set()


_TEAMS_NOTIFY_CENTERS = {"강남센터", "강동센터", "강북센터", "강서센터"}


def _check_uploaded_out_centers(upload_date: date) -> set:
    """해당 날짜 출고 업로드된 to_center 집합 반환."""
    try:
        res = get_supabase().table(TABLE).select("to_center") \
            .eq("upload_date", upload_date.isoformat()) \
            .eq("direction", "out").execute()
        return {r["to_center"] for r in (res.data or [])}
    except Exception:
        return set()


def _build_teams_card(pivot: pd.DataFrame, ref_date: date) -> dict:
    """센터별 출고 현황 Adaptive Card JSON 생성 (Table 요소 사용)."""
    date_str = f"{ref_date.month}월 {ref_date.day}일"
    cols = [c for c in _COL_ORDER if c in pivot.columns]

    def _cell(text, bold=False, align="Center", color=None, style=None):
        tb = {"type": "TextBlock", "text": str(text) if text else " ",
              "size": "Small", "horizontalAlignment": align, "wrap": False}
        if bold:  tb["weight"] = "Bolder"
        if color: tb["color"]  = color
        cell = {"type": "TableCell", "items": [tb]}
        if style: cell["style"] = style
        return cell

    # 헤더 행
    header_cells = [_cell("종류", bold=True, align="Left", color="Accent", style="accent")]
    for c in cols:
        header_cells.append(_cell(c, bold=True, color="Accent", style="accent"))

    # 데이터 행
    data_rows = []
    for row_key, row_data in pivot.iterrows():
        cells = [_cell(str(row_key), bold=True, align="Left", color="Accent")]
        for c in cols:
            v = row_data.get(c, 0)
            cells.append(_cell(str(int(v)) if v > 0 else ""))
        data_rows.append({"type": "TableRow", "cells": cells})

    # 컬럼 너비: 종류 2, 센터별 1
    col_defs = [{"width": 2}] + [{"width": 1}] * len(cols)

    return {
        "type": "AdaptiveCard",
        "version": "1.5",
        "$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
        "body": [
            {"type": "TextBlock", "text": "📤 센터별 출고 현황",
             "size": "Large", "weight": "Bolder", "color": "Accent"},
            {"type": "TextBlock", "text": f"{date_str}  ·  ✅ 4개 센터 업로드 완료",
             "size": "Small", "color": "Good", "spacing": "Small"},
            {
                "type": "Table",
                "firstRowAsHeaders": True,
                "showGridLines": True,
                "gridStyle": "accent",
                "horizontalCellContentAlignment": "Center",
                "columns": col_defs,
                "rows": [
                    {"type": "TableRow", "cells": header_cells}
                ] + data_rows,
                "spacing": "Medium"
            }
        ]
    }


def _send_teams_out_summary(pivot: pd.DataFrame, ref_date: date) -> bool:
    """센터별 출고 현황을 Teams 그룹채팅으로 전송."""
    webhook_url = st.secrets.get("TEAMS_WEBHOOK_URL", "")
    if not webhook_url:
        return False
    try:
        import urllib.request, json as _json
        payload = {"card": _build_teams_card(pivot, ref_date)}
        data = _json.dumps(payload).encode("utf-8")
        req  = urllib.request.Request(
            webhook_url, data=data,
            headers={"Content-Type": "application/json"}
        )
        urllib.request.urlopen(req, timeout=10)
        return True
    except Exception as e:
        st.warning(f"Teams 알림 전송 실패: {e}")
        return False


@st.experimental_dialog("📤 추가 출고 업로드")
def _extra_upload_dialog(sel_date):
    to_c    = st.selectbox("도착 센터", NON_HUB_CENTERS, key="dlg_xout_to")
    mv_date = st.date_input("이동 날짜", value=sel_date, key="dlg_xout_date")
    uploaded = st.file_uploader(
        "엑셀 파일 (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="dlg_xout_file",
        help="단말기 이동신청서 양식. 1·2번 행은 제목으로 간주해 자동 건너뜁니다.",
    )
    if not uploaded:
        return

    df = parse_terminal_excel(uploaded)
    if df is None or df.empty:
        st.warning("데이터를 읽지 못했습니다.")
        return

    auto_col    = find_trcn_col(df)
    all_cols    = df.columns.tolist()
    default_idx = all_cols.index(auto_col) if auto_col in all_cols else 0
    trcn_col    = st.selectbox("단말기ID 컬럼 선택", all_cols,
                               index=default_idx, key="dlg_xout_col")

    df["_trcn"] = df[trcn_col].astype(str).str.strip()
    cls = _apply_classifications(df["_trcn"])
    df  = pd.concat([df, cls], axis=1)

    valid = df[df["_dtype"] != "미분류"].copy()
    inv   = df[df["_dtype"] == "미분류"]

    c1, c2 = st.columns(2)
    c1.metric("✅ 분류 성공", f"{len(valid)}")
    c2.metric("⚠️ 미분류",   f"{len(inv)}")
    if not inv.empty:
        st.warning(f"⚠️ 미분류 {len(inv)}건 — 저장 제외")
    if valid.empty:
        st.warning("분류 가능한 단말기가 없습니다.")
        return

    dups   = check_dups(valid["_trcn"].tolist(), mv_date, "out")
    new_df = valid[~valid["_trcn"].isin(dups)]
    if dups:
        st.warning(f"⚠️ 중복 {len(dups)}건 제외됨")
    if new_df.empty:
        st.error("저장할 데이터가 없습니다 (전부 중복).")
        return

    st.info(f"저장 예정: **{len(new_df)}건** / 자재센터 → {to_c} / {mv_date}")
    _sa, _ca = st.columns(2)
    if _ca.button("취소", use_container_width=True, key="dlg_xout_cancel"):
        st.session_state.pop("_show_extra_upload", None)
        st.rerun()
    if _sa.button("✅ 추가 저장", type="primary", use_container_width=True, key="dlg_xout_save"):
        uid     = str(uuid.uuid4())
        records = [
            {
                "upload_id":   uid,
                "trcn_id":     row["_trcn"],
                "device_type": row["_dtype"],
                "sub_type":    row["_stype"],
                "from_center": "자재센터",
                "to_center":   to_c,
                "direction":   "out",
                "uploaded_by": user["id"],
                "upload_date": mv_date.isoformat(),
                "file_name":   uploaded.name,
                "notes":       None,
            }
            for _, row in new_df.iterrows()
        ]
        if save_terminal(records):
            st.session_state["_extra_upload_done"] = f"✅ {len(records)}건 추가 저장 완료!"
            st.session_state.pop("_show_extra_upload", None)
            st.rerun()


def save_terminal(records: list) -> bool:
    try:
        get_supabase().table(TABLE).insert(records).execute()
        return True
    except Exception as e:
        st.error(f"저장 실패: {e}")
        return False


def delete_by_upload_id(upload_id: str) -> bool:
    try:
        get_supabase().table(TABLE).delete().eq("upload_id", upload_id).execute()
        return True
    except Exception as e:
        st.error(f"삭제 실패: {e}")
        return False


def delete_by_id(record_id: str) -> bool:
    try:
        get_supabase().table(TABLE).delete().eq("id", record_id).execute()
        return True
    except Exception as e:
        st.error(f"삭제 실패: {e}")
        return False


def update_record(record_id: str, device_type: str, sub_type: str) -> bool:
    try:
        get_supabase().table(TABLE).update(
            {"device_type": device_type, "sub_type": sub_type}
        ).eq("id", record_id).execute()
        return True
    except Exception as e:
        st.error(f"수정 실패: {e}")
        return False


# ══════════════════════════════════════════════════════════════════════════════
# 수정/삭제 UI
# ══════════════════════════════════════════════════════════════════════════════

@st.experimental_dialog("단말기 기록 수정")
def _edit_record_dialog(rec: dict):
    new_trcn = st.text_input("IH (TRCN_ID)", value=rec["trcn_id"], key="ed_trcn")
    _di = DEVICE_ORDER.index(rec["device_type"]) if rec["device_type"] in DEVICE_ORDER else 0
    _si = SUB_ORDER.index(rec["sub_type"]) if rec["sub_type"] in SUB_ORDER else 0
    new_dt = st.selectbox("기종", DEVICE_ORDER, index=_di, key="ed_dtype")
    new_st = st.selectbox("유형", SUB_ORDER,    index=_si, key="ed_stype")
    c1, c2 = st.columns(2)
    if c1.button("저장", type="primary", use_container_width=True):
        if not new_trcn.strip():
            st.warning("IH를 입력해 주세요.")
        else:
            try:
                get_supabase().table(TABLE).update({
                    "trcn_id":     new_trcn.strip(),
                    "device_type": new_dt,
                    "sub_type":    new_st,
                }).eq("id", rec["id"]).execute()
                st.success("수정됐습니다.")
                st.rerun()
            except Exception as e:
                st.error(f"수정 실패: {e}")
    if c2.button("취소", use_container_width=True):
        st.rerun()


def render_manage_section(rows: list, direction: str, center_filter: str | None, key_prefix: str):
    """업로드 배치별 수정·삭제 UI."""
    if not rows:
        st.info("수정할 데이터가 없습니다.")
        return
    df = pd.DataFrame(rows)
    if center_filter:
        ctr_col = "to_center" if direction == "out" else "from_center"
        df = df[df[ctr_col] == center_filter]
    if df.empty:
        st.info("수정할 데이터가 없습니다.")
        return

    for upload_id, grp in df.groupby("upload_id"):
        fname      = grp["file_name"].iloc[0] or str(upload_id)[:8]
        upl_time   = str(grp["uploaded_at"].iloc[0])[:16].replace("T", " ")
        cnt        = len(grp)
        with st.container(border=True):
            hc1, hc2 = st.columns([4, 1])
            hc1.markdown(f"**📁 {fname}** &nbsp; `{cnt}건` &nbsp; {upl_time}")
            if hc2.button("🗑️ 전체삭제", key=f"{key_prefix}_batch_{upload_id}",
                          type="secondary", use_container_width=True):
                if delete_by_upload_id(upload_id):
                    st.success(f"{cnt}건 삭제됐습니다.")
                    st.rerun()
            st.divider()
            for _, row in grp.iterrows():
                c1, c2, c3, c4, c5 = st.columns([3, 1, 1, 1, 1])
                c1.caption(f"`{row['trcn_id']}`")
                c2.caption(row["device_type"])
                c3.caption(row["sub_type"])
                if c4.button("✏️", key=f"{key_prefix}_edit_{row['id']}", help="분류 수정"):
                    _edit_record_dialog(row.to_dict())
                if c5.button("🗑️", key=f"{key_prefix}_del_{row['id']}", help="삭제"):
                    if delete_by_id(row["id"]):
                        st.rerun()


# ══════════════════════════════════════════════════════════════════════════════
# 인수인계증 Excel 생성
# ══════════════════════════════════════════════════════════════════════════════

def gen_handover_xlsx(rows: list, from_c: str, to_c: str, mv_date: date, notes: str = "") -> bytes:
    import os, re as _re
    from collections import defaultdict
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.worksheet.pagebreak import Break

    _logo_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "atec_logo.png")

    # ── 스타일 ────────────────────────────────────────────────────────────────
    bold14     = Font(name="맑은 고딕", bold=True, size=14)
    bold11     = Font(name="맑은 고딕", bold=True, size=11)
    bold10     = Font(name="맑은 고딕", bold=True, size=10)
    norm10     = Font(name="맑은 고딕", size=10)
    ca         = Alignment(horizontal="center", vertical="center")
    la         = Alignment(horizontal="left",   vertical="center")
    thin       = Side(style="thin")
    bdr        = Border(left=thin, right=thin, top=thin, bottom=thin)
    gray       = PatternFill("solid", fgColor="D9D9D9")
    lblue      = PatternFill("solid", fgColor="BDD7EE")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    date_str   = mv_date.strftime("%Y-%m-%d")

    _XLSX_ROW_ORDER = [
        ("B800", "표출기"), ("B800", "통합단말기"), ("B800", "승하차"), ("B800", "모뎀"),
        ("B700", "표출기"), ("B700", "통합단말기"), ("B700", "모뎀"),
        ("B710", "표출기"), ("B710", "통합단말기"), ("B710", "승하차"), ("B710", "모뎀"),
        ("B620", "승하차"), ("B620", "운전자"), ("B620", "모뎀"),
        ("한강버스", "통합단말기"), ("한강버스", "표출기"), ("한강버스", "승하차"),
    ]
    _order_map = {pair: i for i, pair in enumerate(_XLSX_ROW_ORDER)}

    # ── 공통 헬퍼 ─────────────────────────────────────────────────────────────
    def _set_widths(ws):
        for col, w in zip("ABCD", [12, 30, 28, 22]):
            ws.column_dimensions[col].width = w

    def _add_logo(ws, anchor_row: int):
        if not os.path.exists(_logo_path):
            return
        img = XLImage(_logo_path)
        img.width = 130; img.height = 44
        ws.row_dimensions[anchor_row].height = 36
        ws.add_image(img, f"D{anchor_row}")

    def _write_header(ws, start_r: int, title: str, _from: str, _to: str, _rows: list):
        ws.merge_cells(f"A{start_r}:D{start_r}")
        c = ws.cell(start_r, 1, title); c.font = bold14; c.alignment = ca
        ws.row_dimensions[start_r].height = 30
        for ci, (lbl, val) in enumerate([("출발센터", _from), ("도착센터", _to)], start=1):
            ws.cell(start_r+1, ci*2-1, lbl).font = bold11
            ws.cell(start_r+1, ci*2-1).alignment = la
            ws.cell(start_r+1, ci*2,   val).font = norm10
            ws.cell(start_r+1, ci*2  ).alignment = la
        ws.cell(start_r+2, 1, "날짜").font    = bold11;  ws.cell(start_r+2, 1).alignment = la
        ws.cell(start_r+2, 2, date_str).font  = norm10;  ws.cell(start_r+2, 2).alignment = la
        ws.cell(start_r+2, 3, "총 수량").font = bold11;  ws.cell(start_r+2, 3).alignment = la
        ws.cell(start_r+2, 4, f"{len(_rows):,}대").font = norm10
        ws.cell(start_r+2, 4).alignment = la

    def _write_device_summary(ws, r: int, _rows: list) -> int:
        ws.merge_cells(f"A{r}:D{r}")
        h = ws.cell(r, 1, "종류별 수량 요약"); h.font = bold11; h.alignment = ca; h.fill = lblue
        r += 1
        for ci, hdr in enumerate(["No", "단말기종류", "유형", "수량"], 1):
            c = ws.cell(r, ci, hdr); c.font = bold11; c.alignment = ca; c.fill = gray; c.border = bdr
        r += 1
        if _rows:
            s_df = (
                pd.DataFrame(_rows)
                .groupby(["device_type", "sub_type"], sort=False)
                .size().reset_index(name="cnt")
            )
            s_df["_ord"] = s_df.apply(
                lambda x: _order_map.get((x["device_type"], x["sub_type"]), len(_XLSX_ROW_ORDER)), axis=1)
            s_df = s_df.sort_values("_ord").drop(columns="_ord").reset_index(drop=True)
            for idx, rec in enumerate(s_df.itertuples(), 1):
                for ci, val in enumerate([idx, rec.device_type, rec.sub_type, rec.cnt], 1):
                    c = ws.cell(r, ci, val); c.font = norm10; c.alignment = ca; c.border = bdr
                r += 1
            ws.merge_cells(f"A{r}:C{r}")
            c = ws.cell(r, 1, "합계"); c.font = bold10; c.alignment = ca; c.fill = total_fill; c.border = bdr
            for col in (2, 3):
                ws.cell(r, col).fill = total_fill; ws.cell(r, col).border = bdr
            c4 = ws.cell(r, 4, len(_rows)); c4.font = bold10; c4.alignment = ca
            c4.fill = total_fill; c4.border = bdr
            r += 1
        return r

    def _write_notes_sig(ws, r: int, _notes: str) -> int:
        r += 1
        ws.merge_cells(f"A{r}:D{r}")
        hh = ws.cell(r, 1, "비고"); hh.font = bold11; hh.alignment = ca; hh.fill = gray; hh.border = bdr
        for col in range(2, 5):
            ws.cell(r, col).fill = gray; ws.cell(r, col).border = bdr
        r += 1
        ns = r
        ws.merge_cells(f"A{ns}:D{ns+3}")
        nc = ws.cell(ns, 1, _notes or ""); nc.font = norm10
        nc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        for ri in range(ns, ns + 4):
            ws.row_dimensions[ri].height = 18
            for ci in range(1, 5):
                t  = thin if ri == ns     else Side(style=None)
                b  = thin if ri == ns + 3 else Side(style=None)
                l  = thin if ci == 1      else Side(style=None)
                rr = thin if ci == 4      else Side(style=None)
                ws.cell(ri, ci).border = Border(left=l, right=rr, top=t, bottom=b)
        r += 4
        r += 2
        for ci, lbl in [(2, "인계자"), (4, "인수자")]:
            ws.cell(r, ci, lbl).font = bold11
        r += 1
        ws.cell(r, 2, "(서명)").alignment = ca
        ws.cell(r, 4, "(서명)").alignment = ca
        r += 1
        return r

    def _write_trcn_list(ws, r: int, _rows: list) -> int:
        CHUNK = 30  # 열당 최대 행 수
        trcn_list = sorted(rec.get("trcn_id", "") for rec in _rows)

        ws.merge_cells(f"A{r}:D{r}")
        h = ws.cell(r, 1, "단말기 IH 목록"); h.font = bold11; h.alignment = ca; h.fill = lblue
        r += 1
        # 헤더: A-B, C-D 각각
        for base in [1, 3]:
            ws.cell(r, base,   "No").font           = bold11; ws.cell(r, base  ).alignment = ca
            ws.cell(r, base  ).fill = gray;                   ws.cell(r, base  ).border    = bdr
            ws.cell(r, base+1, "IH (TRCN_ID)").font = bold11; ws.cell(r, base+1).alignment = ca
            ws.cell(r, base+1).fill = gray;                   ws.cell(r, base+1).border    = bdr
        r += 1
        # CHUNK*2 개씩 블록: 왼쪽(A-B) 30개, 오른쪽(C-D) 30개
        for blk in range(0, max(len(trcn_list), 1), CHUNK * 2):
            left  = trcn_list[blk         : blk + CHUNK]
            right = trcn_list[blk + CHUNK : blk + CHUNK * 2]
            for i in range(max(len(left), len(right))):
                if i < len(left):
                    no = blk + i + 1
                    ws.cell(r, 1, no).font         = norm10; ws.cell(r, 1).alignment = ca; ws.cell(r, 1).border = bdr
                    ws.cell(r, 2, left[i]).font    = norm10; ws.cell(r, 2).alignment = ca; ws.cell(r, 2).border = bdr
                if i < len(right):
                    no = blk + CHUNK + i + 1
                    ws.cell(r, 3, no).font         = norm10; ws.cell(r, 3).alignment = ca; ws.cell(r, 3).border = bdr
                    ws.cell(r, 4, right[i]).font   = norm10; ws.cell(r, 4).alignment = ca; ws.cell(r, 4).border = bdr
                r += 1
        return r

    def _write_center_sheet(ws, _rows: list, _from: str, _to: str, _notes: str = ""):
        _set_widths(ws)
        _write_header(ws, 1, "단말기 이동 인수인계증 — 요약", _from, _to, _rows)
        r = 5
        r = _write_device_summary(ws, r, _rows)
        r = _write_notes_sig(ws, r, _notes)
        r += 1; _add_logo(ws, r); r += 1
        ws.row_breaks.append(Break(id=r)); r += 1
        _write_header(ws, r, "단말기 이동 인수인계증 — 상세", _from, _to, _rows)
        r += 4
        r = _write_trcn_list(ws, r, _rows)
        r += 1; _add_logo(ws, r); r += 1
        ws.print_area                    = f"A1:D{r}"
        ws.page_setup.fitToPage          = True
        ws.page_setup.fitToWidth         = 1
        ws.page_setup.fitToHeight        = 0
        ws.page_setup.horizontalCentered = True
        ws.page_margins.left   = 0.7
        ws.page_margins.right  = 0.7
        ws.page_margins.top    = 0.75
        ws.page_margins.bottom = 0.75

    # ── 멀티/싱글 분기 ────────────────────────────────────────────────────────
    center_field = None
    if   to_c   == "타센터": center_field = "to_center"
    elif from_c == "타센터": center_field = "from_center"

    wb = openpyxl.Workbook()

    if center_field:
        groups: dict[str, list] = defaultdict(list)
        for rec in rows:
            groups[rec.get(center_field) or "미확인"].append(rec)
        sorted_centers = sorted(groups.keys())
        total_cnt = len(rows)

        # ── Sheet 1: 센터별 전체 요약 ─────────────────────────────────────
        from openpyxl.utils import get_column_letter
        _direction  = "out" if center_field == "to_center" else "in"
        _cpivot     = build_center_pivot(rows, direction=_direction)
        _pcols      = list(_cpivot.columns) if _cpivot is not None else []
        _total_cols = len(_pcols) + 2       # 종류 열 + 센터 N개 + 합계 열
        _last_L     = get_column_letter(_total_cols)

        ws0 = wb.active; ws0.title = "전체요약"
        ws0.column_dimensions["A"].width = 16
        for _ci in range(2, _total_cols + 1):
            ws0.column_dimensions[get_column_letter(_ci)].width = 9

        _lbl_from = from_c if from_c != "타센터" else "각 센터"
        _lbl_to   = to_c   if to_c   != "타센터" else "각 센터"

        ws0.merge_cells(f"A1:{_last_L}1")
        c = ws0.cell(1, 1, "단말기 이동 인수인계증 — 센터별 요약"); c.font = bold14; c.alignment = ca
        ws0.row_dimensions[1].height = 30
        for ci, (lbl, val) in enumerate([("출발센터", _lbl_from), ("도착센터", _lbl_to)], start=1):
            ws0.cell(2, ci*2-1, lbl).font = bold11; ws0.cell(2, ci*2-1).alignment = la
            ws0.cell(2, ci*2,   val).font = norm10; ws0.cell(2, ci*2  ).alignment = la
        ws0.cell(3, 1, "날짜").font    = bold11; ws0.cell(3, 1).alignment = la
        ws0.cell(3, 2, date_str).font  = norm10; ws0.cell(3, 2).alignment = la
        ws0.cell(3, 3, "총 수량").font = bold11; ws0.cell(3, 3).alignment = la
        ws0.cell(3, 4, f"{total_cnt:,}대").font = norm10; ws0.cell(3, 4).alignment = la

        r = 5
        _sec_label = "센터별 출고 현황" if _direction == "out" else "센터별 입고 현황"
        ws0.merge_cells(f"A{r}:{_last_L}{r}")
        h = ws0.cell(r, 1, _sec_label); h.font = bold11; h.alignment = ca; h.fill = lblue
        r += 1

        if _cpivot is not None:
            # 컬럼 헤더: 종류 | 센터1 | 센터2 | … | 합계
            _sum_ci = len(_pcols) + 2
            ws0.cell(r, 1, "종류").font = bold11; ws0.cell(r, 1).alignment = ca
            ws0.cell(r, 1).fill = gray; ws0.cell(r, 1).border = bdr
            for _i, _cn in enumerate(_pcols, 2):
                c = ws0.cell(r, _i, _cn); c.font = bold11; c.alignment = ca; c.fill = gray; c.border = bdr
            c = ws0.cell(r, _sum_ci, "합계"); c.font = bold11; c.alignment = ca; c.fill = gray; c.border = bdr
            r += 1

            _col_totals = [0] * len(_pcols)
            for _rk, _rd in _cpivot.iterrows():
                ws0.cell(r, 1, _rk).font = norm10; ws0.cell(r, 1).alignment = la; ws0.cell(r, 1).border = bdr
                _row_total = 0
                for _i, _cn in enumerate(_pcols, 2):
                    _v = int(_rd.get(_cn, 0))
                    c = ws0.cell(r, _i, _v if _v > 0 else ""); c.font = norm10; c.alignment = ca; c.border = bdr
                    _row_total += _v; _col_totals[_i - 2] += _v
                c = ws0.cell(r, _sum_ci, _row_total); c.font = bold10; c.alignment = ca
                c.fill = total_fill; c.border = bdr
                r += 1

            # 합계 행
            ws0.cell(r, 1, "합계").font = bold10; ws0.cell(r, 1).alignment = ca
            ws0.cell(r, 1).fill = total_fill; ws0.cell(r, 1).border = bdr
            for _i, _ct in enumerate(_col_totals, 2):
                c = ws0.cell(r, _i, _ct); c.font = bold10; c.alignment = ca
                c.fill = total_fill; c.border = bdr
            c = ws0.cell(r, _sum_ci, total_cnt); c.font = bold10; c.alignment = ca
            c.fill = total_fill; c.border = bdr
        else:
            ws0.merge_cells(f"A{r}:{_last_L}{r}")
            ws0.cell(r, 1, "데이터 없음").font = norm10; ws0.cell(r, 1).alignment = ca

        # ── 센터별 시트 ───────────────────────────────────────────────────
        for cname in sorted_centers:
            safe = _re.sub(r'[\\/*?:\[\]]', '_', cname)[:31]
            ws_c = wb.create_sheet(title=safe)
            _from = from_c if from_c != "타센터" else cname
            _to   = to_c   if to_c   != "타센터" else cname
            _write_center_sheet(ws_c, groups[cname], _from, _to, notes)
    else:
        ws = wb.active; ws.title = "인수인계증"
        _write_center_sheet(ws, rows, from_c, to_c, notes)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# 업로드 처리 (공통 로직)
# ══════════════════════════════════════════════════════════════════════════════

def _upload_section(direction: str, from_c: str, to_c_fixed: str | None, key_prefix: str):
    """direction='out' or 'in'. to_c_fixed=None → 선택박스."""
    dir_label = "출고" if direction == "out" else "입고"

    if to_c_fixed is None:
        to_c = st.selectbox("도착 센터", NON_HUB_CENTERS, key=f"{key_prefix}_to")
    else:
        to_c = to_c_fixed
        st.markdown(f"**도착 센터:** `{to_c}`")

    mv_date = st.date_input("이동 날짜", value=_today_kst(), key=f"{key_prefix}_date")
    uploaded = st.file_uploader(
        "엑셀 파일 (.xls / .xlsx, 3번 행 = 헤더)",
        type=["xls", "xlsx"],
        key=f"{key_prefix}_file",
        help="단말기 이동신청서 양식. 1·2번 행은 제목으로 간주해 자동 건너뜁니다.",
    )

    if not uploaded:
        return

    df = parse_terminal_excel(uploaded)
    if df is None or df.empty:
        st.warning("데이터를 읽지 못했습니다.")
        return

    auto_col = find_trcn_col(df)
    all_cols = df.columns.tolist()
    default_idx = all_cols.index(auto_col) if auto_col in all_cols else 0
    trcn_col = st.selectbox(
        "단말기ID 컬럼 선택",
        all_cols,
        index=default_idx,
        key=f"{key_prefix}_col",
    )

    df["_trcn"] = df[trcn_col].astype(str).str.strip()
    cls = _apply_classifications(df["_trcn"])
    df = pd.concat([df, cls], axis=1)

    valid = df[df["_dtype"] != "미분류"].copy()
    inv   = df[df["_dtype"] == "미분류"]

    c1, c2 = st.columns(2)
    c1.metric("✅ 분류 성공", f"{len(valid)}")
    c2.metric("⚠️ 미분류", f"{len(inv)}")

    if not inv.empty:
        st.warning(f"⚠️ 미분류 {len(inv)}건 — 규칙 미매칭, 저장 제외")
        inv_show = inv[[trcn_col, "_trcn"]].copy()
        inv_show.columns = ["원본값", "추출값(숫자)"]
        inv_show["자릿수"] = inv_show["추출값(숫자)"].str.len()
        st.dataframe(inv_show, use_container_width=True, hide_index=True)

    if valid.empty:
        st.warning("분류 가능한 단말기가 없습니다. 컬럼을 확인하세요.")
        return

    s_df = (
        valid.groupby(["_dtype", "_stype"]).size()
        .reset_index(name="수량")
        .rename(columns={"_dtype": "종류", "_stype": "유형"})
    )
    st.caption("분류 요약")
    st.dataframe(s_df, use_container_width=True, hide_index=True)

    dups   = check_dups(valid["_trcn"].tolist(), mv_date, direction)
    dup_df = valid[valid["_trcn"].isin(dups)]
    new_df = valid[~valid["_trcn"].isin(dups)]

    if not dup_df.empty:
        st.warning(f"⚠️ {len(dup_df)}건이 같은 날짜·방향으로 이미 DB에 존재합니다.")
        dup_show = dup_df[[trcn_col, "_trcn", "_dtype", "_stype"]].copy()
        dup_show.columns = ["원본값", "TRCN_ID", "기종", "유형"]
        st.dataframe(dup_show, use_container_width=True, hide_index=True)
        force = st.checkbox("중복 무시하고 강제 저장", key=f"{key_prefix}_force")
        if force:
            new_df = valid.copy()

    if new_df.empty:
        st.error("저장할 데이터가 없습니다.")
        return

    st.info(f"저장 예정: **{len(new_df)}건** / {from_c} → {to_c} / {mv_date}")
    notes = st.text_area("비고 (선택)", placeholder="인수인계증에 표시될 메모를 입력하세요.", key=f"{key_prefix}_notes", height=80)

    if st.button(f"✅ {dir_label} 데이터 저장", type="primary", key=f"{key_prefix}_save"):
        uid = str(uuid.uuid4())
        records = [
            {
                "upload_id":   uid,
                "trcn_id":     row["_trcn"],
                "device_type": row["_dtype"],
                "sub_type":    row["_stype"],
                "from_center": from_c,
                "to_center":   to_c,
                "direction":   direction,
                "uploaded_by": user["id"],
                "upload_date": mv_date.isoformat(),
                "file_name":   uploaded.name,
                "notes":       notes.strip() or None,
            }
            for _, row in new_df.iterrows()
        ]
        if save_terminal(records):
            st.success(f"✅ {len(records)}건 저장 완료!")
            st.rerun()


# ══════════════════════════════════════════════════════════════════════════════
# 사이드바
# ══════════════════════════════════════════════════════════════════════════════
with st.sidebar:
    render_sidebar_header()
    if st.button("자재현황(전체)", use_container_width=True):
        st.switch_page("pages/10_dashboard.py")
    if st.button("버스단말기 현황", use_container_width=True, type="primary"):
        st.session_state.pop("_show_extra_upload", None)
    st.divider()

    render_sidebar_section("재고 관리")
    if st.button("🗂️ 통합 뷰", use_container_width=True):
        st.switch_page("pages/15_combined.py")
    if st.button("📋 입출고 이력", use_container_width=True):
        st.switch_page("pages/04_history.py")
    if not is_role("guest"):
        if st.button("📊 사용내역", use_container_width=True):
            st.switch_page("pages/08_usage_history.py")

    if not is_role("guest"):
        render_sidebar_section("요청")
        if st.button("📦 자재요청현황", use_container_width=True, key="sb_mat"):
            st.switch_page("pages/07_material_requests.py")
        if st.button("🛒 구매 요청", use_container_width=True, key="sb_pur"):
            st.switch_page("pages/11_purchase_requests.py")

    if is_role("admin", "materials"):
        render_sidebar_section("관리")
        if st.button("📍 위치 지도", use_container_width=True):
            st.switch_page("pages/09_rack_map.py")
    if is_role("admin"):
        if st.button("⚙️ 관리자", use_container_width=True):
            st.switch_page("pages/05_admin.py")
        if st.button("🟢 접속 현황", use_container_width=True):
            st.switch_page("pages/13_online_users.py")

    st.divider()
    if st.button("💬 문의하기", use_container_width=True):
        st.switch_page("pages/12_inquiry.py")
    if st.button("👤 마이페이지", use_container_width=True):
        st.switch_page("pages/06_mypage.py")
    if st.button("🚪 로그아웃", use_container_width=True):
        logout()
    render_sidebar_user(user)


# ══════════════════════════════════════════════════════════════════════════════
# 메인 컨텐츠
# ══════════════════════════════════════════════════════════════════════════════
render_top_bar("대시보드(버스단말기)", user)
st.markdown("## 📟 단말기 이동 현황 대시보드")
st.divider()

# 테이블 존재 확인
if not _table_exists():
    st.error("⚠️ Supabase 테이블 `terminal_movements` 가 없습니다. 아래 SQL을 실행하세요.")
    st.code(_SQL_SETUP, language="sql")
    st.stop()

_tab_labels = ["📊 오늘의 현황", "📋 이력 조회", "📄 인수인계증"]
if _is_admin:
    _tab_labels.append("⚙️ 관리")
_tabs = st.tabs(_tab_labels)
tab_dash, tab_hist, tab_cert = _tabs[0], _tabs[1], _tabs[2]
tab_admin = _tabs[3] if _is_admin else None


# ══ Tab 1: 오늘의 현황 ════════════════════════════════════════════════════════
with tab_dash:
    today = _today_kst()
    _dc, _rc = st.columns([6, 1])
    sel_date = _dc.date_input("조회 날짜", value=today, key="dash_date")
    _rc.markdown("<div style='height:27px'></div>", unsafe_allow_html=True)
    if _rc.button("새로고침", key="t1_ref", use_container_width=True):
        st.rerun()

    out_rows = fetch_terminal(direction="out", upload_date=sel_date)
    in_rows  = fetch_terminal(direction="in",  upload_date=sel_date)

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("📤 출고", f"{len(out_rows):,}대")
    k2.metric("📥 입고", f"{len(in_rows):,}대")
    k3.metric("📅 날짜", sel_date.strftime("%Y-%m-%d"))
    k4.metric("👤 소속", user_center)
    st.divider()

    # 추가출고 완료 팝업
    if st.session_state.get("_extra_upload_done"):
        st.success(st.session_state.pop("_extra_upload_done"))
    if st.session_state.get("_show_extra_upload"):
        _extra_upload_dialog(st.session_state["_show_extra_upload"])

    # ── 단말기종류 × 센터 크로스표 (출고 | 입고) ─────────────────────────────
    _tbl_out, _tbl_in = st.columns(2)
    with _tbl_out:
        _out_hdr, _out_extra_btn, _out_teams_btn = st.columns([4, 1, 1])
        _out_hdr.markdown("#### 📤 센터별 출고 현황")
        _cp_out = build_center_pivot(out_rows, direction="out")
        if _is_admin or _is_jjae:
            if _out_extra_btn.button("➕ 추가출고", key="extra_upload_btn",
                                     use_container_width=True, help="출고 데이터 추가 업로드"):
                st.session_state["_show_extra_upload"] = sel_date
                st.rerun()
        if _cp_out is not None:
            render_center_table(_cp_out, sel_date)
            if _is_admin or _is_jjae:
                if _out_teams_btn.button("📨 Teams", key="teams_send_btn",
                                         use_container_width=True, help="Teams 채팅방으로 출고 현황 전송"):
                    if _send_teams_out_summary(_cp_out, sel_date):
                        st.success("📨 Teams 채팅방으로 출고 현황을 전송했습니다.")
                    else:
                        st.error("전송 실패 — secrets.toml의 TEAMS_WEBHOOK_URL을 확인하세요.")
        else:
            st.info("📭 해당 날짜 출고 데이터가 없습니다.")
    with _tbl_in:
        st.markdown("#### 📥 센터별 입고 현황")
        _cp_in = build_center_pivot(in_rows, direction="in")
        if _cp_in is not None:
            render_center_table(_cp_in, sel_date)
        else:
            st.info("📭 해당 날짜 입고 데이터가 없습니다.")
    st.divider()

    left_col, right_col = st.columns(2)

    # ── 출고 현황 ──────────────────────────────────────────────────────────
    with left_col:
        st.markdown("#### 📤 출고 현황 &nbsp; `자재센터 → 타센터`")
        if out_rows:
            st.dataframe(build_pivot(out_rows).reset_index(), use_container_width=True, hide_index=True)
            out_df  = pd.DataFrame(out_rows)
            ctr_out = out_df.groupby("to_center").size().reset_index(name="수량")
            ctr_out.columns = ["도착 센터", "수량"]
            st.caption("▸ 센터별")
            st.dataframe(ctr_out, use_container_width=True, hide_index=True)
        else:
            st.info("📭 오늘 출고 데이터가 없습니다.")

        if _can_up_out:
            with st.expander("📤 출고 데이터 업로드", expanded=False):
                _upload_section("out", "자재센터", None, "out")
            with st.expander("✏️ 출고 데이터 수정/삭제", expanded=False):
                render_manage_section(out_rows, "out", None, "mgout")

    # ── 입고 현황 ──────────────────────────────────────────────────────────
    with right_col:
        st.markdown("#### 📥 입고 현황 &nbsp; `타센터 → 자재센터`")
        if in_rows:
            st.dataframe(build_pivot(in_rows).reset_index(), use_container_width=True, hide_index=True)
            in_df   = pd.DataFrame(in_rows)
            ctr_in  = in_df.groupby("from_center").size().reset_index(name="수량")
            ctr_in.columns = ["출발 센터", "수량"]
            st.caption("▸ 센터별")
            st.dataframe(ctr_in, use_container_width=True, hide_index=True)
        else:
            st.info("📭 오늘 입고 데이터가 없습니다.")

        if _can_up_in:
            with st.expander("📥 입고 데이터 업로드", expanded=False):
                if _can_sel_in:
                    _from_fixed = st.selectbox("출발 센터", NON_HUB_CENTERS, key="in_from_sel")
                else:
                    _from_fixed = user_center
                    st.markdown(f"**출발 센터:** `{user_center}`")
                _upload_section("in", _from_fixed, "자재센터", "in")
            with st.expander("✏️ 입고 데이터 수정/삭제", expanded=False):
                _in_cf = None if _can_sel_in else user_center
                render_manage_section(in_rows, "in", _in_cf, "mgin")


# ══ Tab 2: 이력 조회 ══════════════════════════════════════════════════════════
with tab_hist:
    f1, f2, f3 = st.columns([2, 2, 3])
    h_from  = f1.date_input("시작일", value=_today_kst() - timedelta(days=30), key="h_from")
    h_to    = f2.date_input("종료일", value=_today_kst(),                       key="h_to")
    h_dir   = f3.selectbox("방향", ["전체", "출고 (자재→센터)", "입고 (센터→자재)"],  key="h_dir")
    if st.button("🔄 새로고침", key="h_ref", use_container_width=True):
        st.rerun()

    h_search = st.text_input(
        "검색", placeholder="TRCN_ID · 센터명 · 종류 · 유형 검색",
        key="h_search", label_visibility="collapsed",
    )

    _dir_map = {"전체": None, "출고 (자재→센터)": "out", "입고 (센터→자재)": "in"}
    hist_rows = fetch_terminal(direction=_dir_map[h_dir], date_from=h_from, date_to=h_to)

    if not hist_rows:
        st.info("조회 결과가 없습니다.")
    else:
        h_df = pd.DataFrame(hist_rows)
        h_df["방향"] = h_df["direction"].map({"out": "출고", "in": "입고"})
        show = h_df.rename(columns={
            "upload_date": "이동날짜", "from_center": "출발센터",
            "to_center": "도착센터", "device_type": "종류",
            "sub_type": "유형", "trcn_id": "TRCN_ID", "file_name": "파일명",
        })[["이동날짜", "방향", "출발센터", "도착센터", "종류", "유형", "TRCN_ID", "파일명"]]

        if h_search and h_search.strip():
            _kw = h_search.strip().lower()
            _mask = show.apply(
                lambda col: col.astype(str).str.lower().str.contains(_kw, na=False)
            ).any(axis=1)
            show = show[_mask]

        st.markdown(f"**총 {len(show):,}건**")
        st.dataframe(show, use_container_width=True, hide_index=True, height=480)

        xbuf = io.BytesIO()
        with pd.ExcelWriter(xbuf, engine="openpyxl") as xw:
            show.to_excel(xw, index=False, sheet_name="이력")
        st.download_button(
            "📥 Excel 다운로드",
            data=xbuf.getvalue(),
            file_name=f"단말기이력_{h_from}_{h_to}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="hist_dl",
        )


# ══ Tab 3: 인수인계증 ══════════════════════════════════════════════════════════
with tab_cert:
    st.markdown("#### 📄 인수인계증 생성")
    c1, c2, c3 = st.columns(3)
    cert_date   = c1.date_input("이동 날짜", value=_today_kst(), key="cert_date")
    _cert_dir_default = 0 if _is_jjae else 1
    _cert_center_opts = ["전체"] + NON_HUB_CENTERS
    _cert_center_default = 0 if _is_jjae else (
        _cert_center_opts.index(user_center) if user_center in _cert_center_opts else 0
    )
    cert_dir    = c2.selectbox("방향", ["출고 (자재→센터)", "입고 (센터→자재)"], index=_cert_dir_default, key="cert_dir")
    cert_center = c3.selectbox("센터 필터 (선택)", _cert_center_opts, index=_cert_center_default, key="cert_center")

    if st.button("🔍 검색", key="cert_go"):
        _cv = "out" if "출고" in cert_dir else "in"
        c_rows = fetch_terminal(direction=_cv, upload_date=cert_date)
        if cert_center != "전체":
            c_rows = [r for r in c_rows if (
                r["to_center"] == cert_center if _cv == "out" else r["from_center"] == cert_center
            )]
        st.session_state.update({
            "cert_rows": c_rows, "cert_dv": _cv,
            "cert_dt": cert_date, "cert_cc": cert_center,
        })

    if st.session_state.get("cert_rows") is not None:
        c_rows = st.session_state["cert_rows"]
        _cv    = st.session_state["cert_dv"]
        _dt    = st.session_state["cert_dt"]
        _cc    = st.session_state["cert_cc"]

        _from_c = "자재센터" if _cv == "out" else (_cc if _cc != "전체" else "타센터")
        _to_c   = (_cc if _cc != "전체" else "타센터") if _cv == "out" else "자재센터"

        if not c_rows:
            st.warning("해당 조건의 데이터가 없습니다.")
        else:
            st.markdown(
                f"**{_from_c} → {_to_c}** &nbsp;|&nbsp; {_dt} &nbsp;|&nbsp; 총 **{len(c_rows):,}대**"
            )
            st.dataframe(build_pivot(c_rows).reset_index(), use_container_width=True, hide_index=True)

            _notes = c_rows[0].get("notes") or "" if c_rows else ""
            xlsx_data = gen_handover_xlsx(c_rows, _from_c, _to_c, _dt, notes=_notes)
            st.download_button(
                "📥 인수인계증 Excel 다운로드",
                data=xlsx_data,
                file_name=f"인수인계증_{_dt}_{_from_c}→{_to_c}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="cert_dl",
            )


# ══ Tab 4: 관리 (admin 전용) ══════════════════════════════════════════════════
if _is_admin and tab_admin is not None:
    with tab_admin:
        st.markdown("#### 🔄 기존 모뎀 데이터 재분류")
        st.caption(
            "DB에 저장된 모뎀 레코드를 현재 분류 규칙으로 재검토합니다.  \n"
            "- **100001~100500** → B620  \n"
            "- **10으로 시작하는 나머지 6자리** → B800"
        )

        if st.button("🔍 변경 대상 미리보기", key="reclassify_preview"):
            with st.spinner("모뎀 레코드 조회 중..."):
                try:
                    res = (
                        get_supabase().table(TABLE)
                        .select("id,trcn_id,device_type,sub_type")
                        .eq("sub_type", "모뎀")
                        .limit(10000)
                        .execute()
                    )
                    all_modems = res.data or []
                except Exception as e:
                    st.error(f"조회 실패: {e}")
                    all_modems = []

            changes = []
            for rec in all_modems:
                new_dtype, new_stype = classify_terminal(rec["trcn_id"])
                if new_dtype != rec["device_type"] or new_stype != rec["sub_type"]:
                    changes.append({
                        "id": rec["id"],
                        "trcn_id": rec["trcn_id"],
                        "기존 기종": rec["device_type"],
                        "변경 기종": new_dtype,
                    })

            st.session_state["reclassify_changes"] = changes
            if not changes:
                st.success("변경이 필요한 레코드가 없습니다.")
            else:
                st.warning(f"변경 대상 **{len(changes)}건**")
                st.dataframe(
                    [{k: v for k, v in c.items() if k != "id"} for c in changes],
                    use_container_width=True, hide_index=True,
                )

        changes = st.session_state.get("reclassify_changes", [])
        if changes:
            if st.button(f"✅ {len(changes)}건 일괄 수정 실행", type="primary", key="reclassify_run"):
                failed = 0
                for i in range(0, len(changes), 100):
                    batch = changes[i:i + 100]
                    ids = [c["id"] for c in batch]
                    dtype_set = set(c["변경 기종"] for c in batch)
                    if len(dtype_set) == 1:
                        try:
                            get_supabase().table(TABLE).update(
                                {"device_type": dtype_set.pop()}
                            ).in_("id", ids).execute()
                        except Exception:
                            failed += len(batch)
                    else:
                        for c in batch:
                            try:
                                get_supabase().table(TABLE).update(
                                    {"device_type": c["변경 기종"]}
                                ).eq("id", c["id"]).execute()
                            except Exception:
                                failed += 1

                if failed:
                    st.error(f"{failed}건 업데이트 실패")
                else:
                    st.success(f"✅ {len(changes)}건 재분류 완료!")
                    st.session_state.pop("reclassify_changes", None)
                    st.rerun()
